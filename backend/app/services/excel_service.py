import io
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from app.schemas.product import VietfulProductInput

def parse_excel_products(file_bytes: bytes) -> List[VietfulProductInput]:
    """
    Parses an Excel (.xlsx) file in-memory and converts rows into VietfulProductInput instances.
    Handles flexible column naming (case-insensitive, snake_case or camelCase).
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet = wb.active
    if sheet is None:
        return []

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    # Identify header columns
    raw_headers = rows[0]
    col_map: Dict[str, int] = {}
    for idx, header in enumerate(raw_headers):
        if header and isinstance(header, str):
            clean_h = header.strip().lower().replace(" ", "").replace("_", "")
            col_map[clean_h] = idx

    # Map standard fields to column indices
    def get_col(field_names: List[str]) -> Optional[int]:
        for fn in field_names:
            clean = fn.lower().replace(" ", "").replace("_", "")
            if clean in col_map:
                return col_map[clean]
        return None

    sku_idx = get_col(["sku", "productsku", "id"])
    name_idx = get_col(["productname", "name", "title"])
    partner_sku_idx = get_col(["partnersku", "partner_sku"])
    price_idx = get_col(["price", "cost", "unitprice"])
    stock_idx = get_col(["stockquantity", "stock", "quantity", "qty"])
    asset_idx = get_col(["assettype", "type"])
    barcode_idx = get_col(["barcodes", "barcode"])

    if sku_idx is None or name_idx is None:
        raise ValueError("Excel file must contain at least 'SKU' and 'ProductName' columns")

    products: List[VietfulProductInput] = []
    for row in rows[1:]:
        if not row or all(v is None for v in row):
            continue

        sku_val = row[sku_idx]
        if sku_val is None or str(sku_val).strip() == "":
            continue  # Skip rows without SKU

        name_val = row[name_idx]
        if name_val is None or str(name_val).strip() == "":
            name_val = f"Product-{sku_val}"

        price_val = 0.0
        if price_idx is not None and row[price_idx] is not None:
            try:
                price_val = float(row[price_idx])
            except (ValueError, TypeError):
                price_val = 0.0

        stock_val = 0
        if stock_idx is not None and row[stock_idx] is not None:
            try:
                stock_val = int(row[stock_idx])
            except (ValueError, TypeError):
                stock_val = 0

        partner_val = None
        if partner_sku_idx is not None and row[partner_sku_idx] is not None:
            partner_val = str(row[partner_sku_idx]).strip()

        asset_val = "Single"
        if asset_idx is not None and row[asset_idx] is not None:
            asset_val = str(row[asset_idx]).strip()

        barcodes_list = []
        if barcode_idx is not None and row[barcode_idx] is not None:
            raw_b = str(row[barcode_idx])
            barcodes_list = [b.strip() for b in raw_b.split(",") if b.strip()]

        product = VietfulProductInput(
            sku=str(sku_val).strip(),
            partnerSKU=partner_val or f"PTR-{sku_val}",
            productName=str(name_val).strip(),
            assetType=asset_val,
            price=price_val,
            stockQuantity=stock_val,
            barcodes=barcodes_list,
            attributes={"source": "EXCEL_UPLOAD"}
        )
        products.append(product)

    return products


def generate_sample_excel() -> bytes:
    """
    Generates a professionally formatted sample Excel spreadsheet with standard columns.
    Returns bytes in-memory without saving to disk.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products Catalog"

    headers = ["SKU", "PartnerSKU", "ProductName", "Price", "StockQuantity", "AssetType", "Barcodes"]
    ws.append(headers)

    # Style header
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    sample_data = [
        ["VF-EXC-001", "PTR-VF-EXC-001", "Chuột không dây Logitech MX Master 3S", 2490000, 85, "Single", "8938501230012"],
        ["VF-EXC-002", "PTR-VF-EXC-002", "Bàn phím cơ Keychron K2 Pro Wireless", 1850000, 42, "Single", "8938501230029"],
        ["VF-EXC-003", "PTR-VF-EXC-003", "Màn hình Dell UltraSharp U2723QE 4K", 12500000, 18, "Bundle", "8938501230036"],
        ["VF-EXC-004", "PTR-VF-EXC-004", "Tai nghe Sony WH-1000XM5 Noise-Canceling", 6990000, 30, "Single", "8938501230043"],
        ["VF-EXC-005", "PTR-VF-EXC-005", "Ổ cứng di động SSD Samsung T7 Shield 1TB", 2790000, 95, "Single", "8938501230050"],
    ]

    for row in sample_data:
        ws.append(row)

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
